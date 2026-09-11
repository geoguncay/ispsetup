"""
Export a Excel de los reportes (Fase 4.3), con openpyxl.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.reports import ClientsReport, ConsumptionReport, OverdueReport, RevenueReport

_HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14)


def _write_table(ws: Worksheet, start_row: int, title: str, headers: list[str], rows: list[list]) -> int:
    """Escribe un título + tabla con encabezado a partir de `start_row`. Devuelve la siguiente fila libre."""
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=12)
    header_row = start_row + 1
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for r, row_values in enumerate(rows, start=header_row + 1):
        for col, value in enumerate(row_values, start=1):
            ws.cell(row=r, column=col, value=value)
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    return header_row + len(rows) + 3


def _buffer_from(wb: Workbook) -> BytesIO:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_revenue_excel(report: RevenueReport) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingresos"
    ws.cell(row=1, column=1, value="Reporte de ingresos").font = _TITLE_FONT
    ws.cell(row=2, column=1, value=f"Total: ${report.total_amount:,.2f} ({report.total_payments} pagos)")

    row = _write_table(
        ws, 4, f"Por período ({report.group_by})", ["Período", "Monto", "Pagos"],
        [[p.label, float(p.amount), p.payments_count] for p in report.by_period],
    )
    row = _write_table(
        ws, row, "Por plan", ["Plan", "Monto", "Pagos"],
        [[p.plan_name, float(p.amount), p.payments_count] for p in report.by_plan],
    )
    _write_table(
        ws, row, "Por sitio", ["Sitio", "Monto", "Pagos"],
        [[s.site_name, float(s.amount), s.payments_count] for s in report.by_site],
    )
    return _buffer_from(wb)


def generate_clients_excel(report: ClientsReport) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.cell(row=1, column=1, value="Reporte de clientes").font = _TITLE_FONT
    ws.cell(
        row=2, column=1,
        value=f"Total: {report.total_clients} — Activos: {report.active_clients} — Suspendidos: {report.suspended_clients}",
    )
    _write_table(
        ws, 4, "Evolución mensual",
        ["Mes", "Nuevos", "Suspensiones", "Bajas (planes cancelados)"],
        [[e.label, e.new_clients, e.suspended_events, e.churned_clients] for e in report.evolution],
    )
    return _buffer_from(wb)


def generate_consumption_excel(report: ConsumptionReport) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Consumo"
    ws.cell(row=1, column=1, value="Reporte de consumo").font = _TITLE_FONT

    row = _write_table(
        ws, 3, "Top consumidores", ["Cliente", "Plan", "Bytes totales"],
        [[c.client_name, c.plan_name or "—", c.total_bytes] for c in report.top_consumers],
    )
    row = _write_table(
        ws, row, "Promedio por plan", ["Plan", "Promedio por cliente (bytes)", "Clientes"],
        [[p.plan_name, p.avg_bytes_per_client, p.clients_count] for p in report.by_plan],
    )
    _write_table(
        ws, row, "Horas pico", ["Hora", "Bytes totales"],
        [[f"{h.hour:02d}:00", h.total_bytes] for h in report.peak_hours],
    )
    return _buffer_from(wb)


def generate_overdue_excel(report: OverdueReport) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Mora"
    ws.cell(row=1, column=1, value="Reporte de mora").font = _TITLE_FONT
    ws.cell(
        row=2, column=1,
        value=f"{report.total_clients} clientes — {report.total_invoices} facturas — ${report.total_amount:,.2f}",
    )
    _write_table(
        ws, 4, "Facturas vencidas",
        ["Cliente", "Plan", "Período", "Vencimiento", "Días de mora", "Monto"],
        [
            [i.client_name, i.plan_name or "—", i.period, i.due_date.strftime("%d/%m/%Y"), i.days_overdue, float(i.amount)]
            for i in report.items
        ],
    )
    return _buffer_from(wb)
